"""팀별 손익계산서 파일 파싱.

.docs/03_데이터정제.md §6.2에서 확인된 구조 두 가지를 모두 지원한다(파일마다 다름,
추정하지 않고 실측된 형태를 모두 인식한다):

1. **가로형(더미 샘플로 최초 확인)** — 팀별 시트(모티브/고정형/대리점/차량oe)로
   나뉘고, 각 시트는 계정과목(행) × 월(열) 구조다. 매출원가(A)Tot·매출총이익(A)·
   판관비(Total)·영업이익(A) 4개 소계 행은 원본이 각각 `=SUM(B8:B40)`, `=B5-B6`,
   `=SUM(B44:B80)`, `=B41-B42` 수식(열은 월마다 이동)인데 캐시값이 비어 있을 수 있어,
   이 형식은 캐시값을 신뢰하지 않고 확인된 수식 그대로 직접 계산한다. 연도는 "국내"
   (전사 집계) 시트 제목에서 추출한다.
2. **세로형(실제 업로드 형식으로 확인, `손익계산서_더미용.xlsx`)** — 시트 1개에
   행 = 팀×월 조합(예: 모티브 1~12월, 고정형 1~12월, ...), 열 = 81개 계정과목이다.
   이 형식은 **파일 어디에도 연도 정보가 없다** — 사용자 확인 결과, 같은 배치로
   업로드되는 실적 파일의 target_year를 그대로 쓰기로 했다(`parse_team_pl_workbook`의
   `year` 인자, `app/routers/uploads.py`에서 실적 파일 기간 확정 후에 호출한다).
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass

from openpyxl import load_workbook

SUMMARY_SHEET = "국내"

# 시트명 -> 앱 전체에서 쓰는 표준 team 값. 이 파일의 시트명은 앱의 표준 팀 명칭과
# 다르다는 것을 실측으로 확인했다("대리점"→"차량대리점", "차량oe"→"차량OE") — 판매계획
# 파일에서 시트명과 VLOOKUP 리터럴이 서로 달랐던 것과 같은 종류의 불일치이므로,
# 추정하지 않고 표준 팀 명칭으로 정규화해서 저장한다.
TEAM_SHEET_TO_TEAM = {
    "모티브": "모티브",
    "고정형": "고정형",
    "대리점": "차량대리점",
    "차량oe": "차량OE",
}

# 헤더 행("구분") 기준 상대 행 번호 — .docs/03_데이터정제.md §6.2에서 4개 팀 시트
# 전부 동일한 것을 실측으로 확인했다(절대 행이 아니라 헤더 행 기준 오프셋으로 둔다).
SALES_AMOUNT_OFFSET = 2  # 매출액(Total)
COGS_TOTAL_OFFSET = 3  # 매출원가(A)Tot = SUM(원가 세부 항목)
COGS_DETAIL_START_OFFSET = 5
COGS_DETAIL_END_OFFSET = 37
GROSS_PROFIT_OFFSET = 38  # 매출총이익(A) = 매출액 - 매출원가
SGA_TOTAL_OFFSET = 39  # 판관비(Total) = SUM(판관비 세부 항목)
SGA_DETAIL_START_OFFSET = 41
SGA_DETAIL_END_OFFSET = 77
OPERATING_PROFIT_OFFSET = 78  # 영업이익(A) = 매출총이익 - 판관비


class TeamPLValidationError(ValueError):
    def __init__(self, message: str):
        super().__init__(message)


@dataclass
class TeamPLEntry:
    team: str
    account_item: str
    year: int
    month: int
    amount: float | None


def _find_year(ws) -> int:
    for row in ws.iter_rows(max_row=10, values_only=True):
        for value in row:
            if isinstance(value, str):
                m = re.search(r"(\d{4})년", value)
                if m:
                    return int(m.group(1))
    raise TeamPLValidationError(
        f"'{ws.title}' 시트 제목에서 연도를 찾을 수 없습니다 (예: '2026년...')."
    )


def _find_header_row(ws) -> int:
    for row_idx in range(1, ws.max_row + 1):
        if ws.cell(row=row_idx, column=1).value == "구분":
            return row_idx
    raise TeamPLValidationError(f"'{ws.title}' 시트에서 헤더 행('구분' 컬럼)을 찾을 수 없습니다.")


def _month_value(ws, row: int, month: int) -> float | None:
    value = ws.cell(row=row, column=1 + month).value
    return float(value) if isinstance(value, (int, float)) else None


def _sum_range(ws, start_row: int, end_row: int, month: int) -> float:
    total = 0.0
    for row in range(start_row, end_row + 1):
        value = _month_value(ws, row, month)
        if value is not None:
            total += value
    return total


def _parse_team_sheet(ws, team: str, year: int) -> list[TeamPLEntry]:
    header_row = _find_header_row(ws)
    entries: list[TeamPLEntry] = []

    for row_idx in range(header_row + 1, ws.max_row + 1):
        account_item = ws.cell(row=row_idx, column=1).value
        if not account_item or not str(account_item).strip():
            continue
        account_item = str(account_item).strip()
        offset = row_idx - header_row

        for month in range(1, 13):
            if offset == COGS_TOTAL_OFFSET:
                amount = _sum_range(
                    ws, header_row + COGS_DETAIL_START_OFFSET, header_row + COGS_DETAIL_END_OFFSET, month
                )
            elif offset == GROSS_PROFIT_OFFSET:
                sales = _month_value(ws, header_row + SALES_AMOUNT_OFFSET, month) or 0.0
                cogs = _sum_range(
                    ws, header_row + COGS_DETAIL_START_OFFSET, header_row + COGS_DETAIL_END_OFFSET, month
                )
                amount = sales - cogs
            elif offset == SGA_TOTAL_OFFSET:
                amount = _sum_range(
                    ws, header_row + SGA_DETAIL_START_OFFSET, header_row + SGA_DETAIL_END_OFFSET, month
                )
            elif offset == OPERATING_PROFIT_OFFSET:
                sales = _month_value(ws, header_row + SALES_AMOUNT_OFFSET, month) or 0.0
                cogs = _sum_range(
                    ws, header_row + COGS_DETAIL_START_OFFSET, header_row + COGS_DETAIL_END_OFFSET, month
                )
                gross_profit = sales - cogs
                sga = _sum_range(
                    ws, header_row + SGA_DETAIL_START_OFFSET, header_row + SGA_DETAIL_END_OFFSET, month
                )
                amount = gross_profit - sga
            else:
                amount = _month_value(ws, row_idx, month)

            entries.append(TeamPLEntry(team=team, account_item=account_item, year=year, month=month, amount=amount))

    return entries


def _is_long_format_sheet(ws) -> bool:
    first_two = [ws.cell(row=1, column=c).value for c in (1, 2)]
    return first_two == ["팀", "월"]


def _parse_long_format(ws, year: int) -> list[TeamPLEntry]:
    account_items = [
        str(ws.cell(row=1, column=c).value).strip()
        for c in range(3, ws.max_column + 1)
        if ws.cell(row=1, column=c).value is not None
    ]
    entries: list[TeamPLEntry] = []
    for row_idx in range(2, ws.max_row + 1):
        team = ws.cell(row=row_idx, column=1).value
        month_raw = ws.cell(row=row_idx, column=2).value
        if not team or not month_raw:
            continue
        month_digits = re.sub(r"\D", "", str(month_raw))
        if not month_digits:
            continue
        month = int(month_digits)

        for col_offset, account_item in enumerate(account_items):
            value = ws.cell(row=row_idx, column=3 + col_offset).value
            amount = float(value) if isinstance(value, (int, float)) else None
            entries.append(
                TeamPLEntry(team=str(team).strip(), account_item=account_item, year=year, month=month, amount=amount)
            )
    return entries


def parse_team_pl_workbook(content: bytes, year: int | None = None) -> list[TeamPLEntry]:
    """팀별 손익계산서 파일을 파싱한다. 가로형(팀별 시트)과 세로형(팀×월 행, 시트 1개)
    두 형식을 모두 인식한다(모듈 docstring 참고). 세로형은 파일에 연도 정보가 없어
    `year`(같은 배치의 실적 파일 target_year)가 필요하다 — 가로형은 이 인자를 무시하고
    "국내" 시트 제목에서 자체적으로 연도를 찾는다.
    """
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:  # noqa: BLE001 — 어떤 형식 오류든 사용자에게 동일하게 안내
        raise TeamPLValidationError(f"손익계산서 파일을 열 수 없습니다: {exc}") from exc

    present_team_sheets = [name for name in TEAM_SHEET_TO_TEAM if name in wb.sheetnames]
    if present_team_sheets:
        if SUMMARY_SHEET not in wb.sheetnames:
            raise TeamPLValidationError(f"손익계산서 파일에 '{SUMMARY_SHEET}' 시트가 없어 연도를 확인할 수 없습니다.")
        found_year = _find_year(wb[SUMMARY_SHEET])
        entries: list[TeamPLEntry] = []
        for sheet_name in present_team_sheets:
            entries.extend(_parse_team_sheet(wb[sheet_name], TEAM_SHEET_TO_TEAM[sheet_name], found_year))
        return entries

    long_format_sheets = [name for name in wb.sheetnames if _is_long_format_sheet(wb[name])]
    if long_format_sheets:
        if year is None:
            raise TeamPLValidationError(
                "세로형 손익계산서 파일은 연도 정보가 없어 실적 파일의 기간을 함께 확인해야 합니다."
            )
        entries = []
        for sheet_name in long_format_sheets:
            entries.extend(_parse_long_format(wb[sheet_name], year))
        return entries

    raise TeamPLValidationError(
        f"손익계산서 파일 형식을 인식할 수 없습니다 (필요: 팀별 시트({list(TEAM_SHEET_TO_TEAM)} 중 최소 1개) "
        "또는 '팀'·'월' 컬럼으로 시작하는 세로형 시트)."
    )
