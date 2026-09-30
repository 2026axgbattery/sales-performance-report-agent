from io import BytesIO

from openpyxl import load_workbook

from tests.conftest import FIXTURES_DIR

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _upload(client, filenames_and_types, mime="text/csv"):
    files = [
        ("files", (name, open(FIXTURES_DIR / name, "rb"), mime))
        for name, _ in filenames_and_types
    ]
    data = {"file_types": [ftype for _, ftype in filenames_and_types]}
    return client.post("/uploads", files=files, data=data)


def _seed_three_months(client):
    r1 = _upload(client, [("actual_prev_year.csv", "실적"), ("mapping_sample.csv", "매핑표")])
    assert r1.status_code == 200, r1.text
    r2 = _upload(client, [("actual_prev_month.csv", "실적")])
    assert r2.status_code == 200, r2.text
    r3 = _upload(client, [("actual_sample.csv", "실적"), ("plan_sample.xlsx", "계획")])
    assert r3.status_code == 200, r3.text
    return r3.json()["batch_id"]


def _load(content: bytes):
    return load_workbook(BytesIO(content))


def test_export_overview_returns_valid_pdf(client):
    # 사용자 확인: Overview 다운로드는 엑셀이 아니라 PDF다.
    batch_id = _seed_three_months(client)
    response = client.get(f"/batches/{batch_id}/export/overview")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert "attachment" in response.headers["content-disposition"]
    assert "filename*=UTF-8''Overview_" in response.headers["content-disposition"]
    assert response.content.startswith(b"%PDF-")


def test_export_anomalies_returns_valid_xlsx(client):
    batch_id = _seed_three_months(client)
    response = client.get(f"/batches/{batch_id}/export/anomalies")
    assert response.status_code == 200

    wb = _load(response.content)
    ws = wb.active
    assert ws.max_row > 1  # 헤더 + 최소 1개 이상징후


def test_export_report_excludes_toggled_items(client):
    batch_id = _seed_three_months(client)
    draft = client.post("/reports", json={"batch_id": batch_id}).json()
    excluded_item = draft["items"][0]
    client.patch(
        f"/reports/{draft['draft_id']}/items/{excluded_item['item_id']}",
        json={"is_excluded": True},
    )

    response = client.get(f"/reports/{draft['draft_id']}/export")
    assert response.status_code == 200
    wb = _load(response.content)
    ws = wb.active
    comments = [row[3].value for row in ws.iter_rows(min_row=2)]
    assert excluded_item["auto_comment"] not in comments
    assert ws.max_row == len(draft["items"])  # 헤더 1 + (전체-1개 제외)


def test_export_report_includes_team_trend_charts(client):
    # 사용자 요청(Phase 22, .docs/phase/phase_22_보고서엑셀차트.md): "엑셀 다운로드 시
    # 내용에 그래프도 같이 들어가게" — F7 화면과 같은 팀별 매출액+영업이익 추이 차트가
    # 엑셀 네이티브 차트로 포함돼야 한다.
    batch_id = _seed_three_months(client)
    draft = client.post("/reports", json={"batch_id": batch_id}).json()
    teams = {item["team"] for item in draft["items"] if not item["is_excluded"] and item["team"]}
    assert teams  # 이 fixture는 최소 1개 팀의 이상징후를 만든다(다른 테스트에서 이미 확인)

    response = client.get(f"/reports/{draft['draft_id']}/export")
    assert response.status_code == 200
    wb = _load(response.content)

    assert "추이 차트" in wb.sheetnames
    assert "차트데이터" in wb.sheetnames
    chart_ws = wb["추이 차트"]
    assert len(chart_ws._charts) == len(teams)  # 팀 1개당 콤보 차트 1개

    data_ws = wb["차트데이터"]
    assert data_ws.sheet_state == "hidden"
    # 첫 팀의 12개월 데이터가 실제로 채워져 있는지(헤더 다음 12행) 확인한다.
    header_row_values = [c.value for c in data_ws[1]]
    assert header_row_values[:3] == ["월", "매출액", "영업이익"]
    months_in_block = [data_ws.cell(row=r, column=1).value for r in range(2, 14)]
    assert months_in_block == list(range(1, 13))


def test_export_unknown_batch_returns_404(client):
    assert client.get("/batches/does-not-exist/export/overview").status_code == 404
    assert client.get("/batches/does-not-exist/export/anomalies").status_code == 404


def test_export_unknown_report_returns_404(client):
    assert client.get("/reports/does-not-exist/export").status_code == 404
    assert client.get("/reports/does-not-exist/export/refined").status_code == 404


def test_export_refined_returns_full_rearrange_columns(client):
    # F7 신규 다운로드 버튼(사용자 요청: "raw와 맵핑을 거쳐 정제된 실적 Re-arrange용
    # 수식 시트 전체, xlsx 양식").
    batch_id = _seed_three_months(client)
    draft = client.post("/reports", json={"batch_id": batch_id}).json()

    response = client.get(f"/reports/{draft['draft_id']}/export/refined")
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE

    wb = _load(response.content)
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    # 정제 결과의 핵심 계정과목이 한글 라벨로 전부 포함돼 있어야 한다(추정 라벨 없이
    # app/services/refinement.py 주석과 03_데이터정제.md에서 실측한 이름 그대로).
    assert "팀" in header
    assert "매출원가(A)Tot" in header
    assert "매출원가(S)Tot" in header
    assert "수량(22.9cell)" in header
    assert ws.max_row > 1  # 헤더 + 배치의 실적 행
